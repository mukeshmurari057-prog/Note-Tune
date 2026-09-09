"""File-to-text extraction for NoteTune."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
import fitz
from docx import Document
import pytesseract

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".jpg", ".jpeg", ".png"}


class ExtractionError(ValueError):
    """Raised when a supported file cannot be read."""


def extract_text(path: str | Path) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type: {suffix or 'unknown'}")

    try:
        if suffix == ".pdf":
            text = _extract_pdf(file_path)
        elif suffix == ".docx":
            text = _extract_docx(file_path)
        elif suffix == ".doc":
            text = _extract_doc(file_path)
        else:
            text = pytesseract.image_to_string(Image.open(file_path))
    except ExtractionError:
        raise
    except (OSError, ValueError, RuntimeError, TypeError) as exc:
        raise ExtractionError(f"Could not extract text from {file_path.name}: {exc}") from exc

    text = clean_text(text)
    if not text:
        raise ExtractionError(f"No readable text found in {file_path.name}")
    return text


def _extract_pdf(path: Path) -> str:
    pages = []
    with fitz.open(path) as document:
        for page in document:
            pages.append(page.get_text("text"))
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(paragraphs)


def _extract_doc(path: Path) -> str:
    """Extract legacy .doc files when the system's antiword utility is available."""
    antiword = shutil.which("antiword")
    if not antiword:
        raise ExtractionError(
            "Legacy .doc files need the antiword utility. Install antiword or save the file as .docx."
        )
    result = subprocess.run(
        [antiword, str(path)],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )
    if result.returncode != 0:
        raise ExtractionError(result.stderr.strip() or "antiword could not read the document")
    return result.stdout


def clean_text(text: str) -> str:
    """Normalize whitespace without removing any words or sentences."""
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def extract_many(files: list[tuple[str, bytes]]) -> str:
    """Extract files in upload order and keep clear source boundaries."""
    sections = []
    with tempfile.TemporaryDirectory() as directory:
        for name, content in files:
            target = Path(directory) / Path(name).name
            target.write_bytes(content)
            sections.append(f"[Source: {Path(name).name}]\n{extract_text(target)}")
    return "\n\n".join(sections)
